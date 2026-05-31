from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/wenjiehu/graduate/模板/ppt模板/ppt学术版模板.pptx")
OUT = ROOT / "docs" / "基于Flink的动态Kafka-Topic发现与订阅机制研究_毕设答辩_模板版.pptx"
NOTES = ROOT / "docs" / "基于Flink的动态Kafka-Topic发现与订阅机制研究_答辩讲稿_模板版.md"
FIG = ROOT / "docs" / "SJTUThesis" / "figures"

FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"

INK = RGBColor(35, 38, 44)
MUTED = RGBColor(98, 103, 111)
RED = RGBColor(143, 31, 36)
BLUE = RGBColor(31, 73, 118)
GREEN = RGBColor(55, 112, 91)
ORANGE = RGBColor(150, 92, 42)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(248, 249, 251)
LIGHT_BLUE = RGBColor(239, 244, 248)
LINE = RGBColor(218, 222, 227)
CODE_BG = RGBColor(250, 250, 250)
CODE_TEXT = RGBColor(48, 54, 61)


def delete_all_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst
    for sld_id in list(xml_slides):
        prs.part.drop_rel(sld_id.rId)
        xml_slides.remove(sld_id)


def set_placeholder_text(slide, idx, text, size=None, bold=None, color=None):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            shape.text = text
            shape.text_frame.word_wrap = True
            shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for p in shape.text_frame.paragraphs:
                p.font.name = FONT_CN
                if size:
                    p.font.size = Pt(size)
                if bold is not None:
                    p.font.bold = bold
                if color:
                    p.font.color.rgb = color
                for run in p.runs:
                    run.font.name = FONT_CN
                    if size:
                        run.font.size = Pt(size)
                    if bold is not None:
                        run.font.bold = bold
                    if color:
                        run.font.color.rgb = color
            return shape
    return None


def add_text(slide, text, x, y, w, h, size=14, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_CN
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align
    return box


def add_title(slide, title, subtitle=None):
    set_placeholder_text(slide, 0, title, 23, True, BLUE)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.58, 8.7, 0.28, 9.6, False, MUTED)


def add_footer(slide, page):
    add_text(slide, f"{page:02d}", 9.45, 0.33, 0.42, 0.18, 8.5, False, MUTED, PP_ALIGN.RIGHT)


def content_panel(slide, x=0.55, y=1.85, w=9.15, h=5.25):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    return shape


def section_label(slide, text, x, y, color=BLUE):
    add_text(slide, text, x, y, 2.5, 0.22, 10.2, True, color)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y + 0.28), Inches(0.58), Inches(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()


def bullet_list(slide, items, x, y, w, h, size=12.8, color=INK, space=3):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT_CN
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(space)
    return box


def numbered_rows(slide, rows, x, y, w, row_h=0.62, colors=None):
    colors = colors or [RED, BLUE]
    for i, (title, body) in enumerate(rows):
        yy = y + i * row_h
        color = colors[i % len(colors)]
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(yy + 0.02), Inches(0.28), Inches(0.28))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.fill.background()
        add_text(slide, str(i + 1), x, yy + 0.055, 0.28, 0.12, 7.2, True, WHITE, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.42, yy, 2.15, 0.24, 11.2, True, BLUE)
        add_text(slide, body, x + 2.15, yy, w - 2.15, 0.28, 10.5, False, INK)


def card(slide, title, body, x, y, w, h, accent=BLUE, size=10.5):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.05
    add_text(slide, title, x + 0.16, y + 0.12, w - 0.28, 0.23, 11.2, True, accent)
    add_text(slide, body, x + 0.16, y + 0.42, w - 0.28, h - 0.48, size, False, INK)


def code_box(slide, title, lines, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = LINE
    add_text(slide, title, x + 0.12, y + 0.1, w - 0.24, 0.2, 9.4, True, BLUE)
    text = "\n".join(lines)
    box = slide.shapes.add_textbox(Inches(x + 0.14), Inches(y + 0.42), Inches(w - 0.28), Inches(h - 0.5))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_CODE
    p.font.size = Pt(8.0)
    p.font.color.rgb = CODE_TEXT
    return box


def metric(slide, value, label, x, y, w=1.85, h=0.82, color=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.06
    add_text(slide, value, x + 0.08, y + 0.1, w - 0.16, 0.26, 16.5, True, color, PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.08, y + 0.47, w - 0.16, 0.22, 7.8, False, MUTED, PP_ALIGN.CENTER)


def add_table(slide, rows, x, y, w, h, font_size=8.4, header_color=WHITE):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_CN
                p.font.size = Pt(font_size)
                p.font.bold = r == 0
                p.font.color.rgb = header_color if r == 0 else INK
    return table


def add_image(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def arrow(slide, x1, y1, x2, y2, color=BLUE):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.4)
    return line


def make_deck():
    prs = Presentation(TEMPLATE)
    delete_all_slides(prs)
    notes = []
    cover = prs.slide_layouts[14]
    contents = prs.slide_layouts[5]
    inner = prs.slide_layouts[2]
    end = prs.slide_layouts[15]

    # 1 Cover
    s = prs.slides.add_slide(cover)
    set_placeholder_text(s, 0, "基于 Flink 的动态 Kafka-Topic\n发现与订阅机制研究", 28, True, BLUE)
    set_placeholder_text(s, 1, "本科毕业设计答辩", 16, False, RED)
    set_placeholder_text(s, 10, "答辩人：胡文杰\n指导教师：任锐\n上海交通大学    2026 年 5 月", 12, False, MUTED)
    notes.append("开场：说明课题研究的是 Flink 写出端，即 Kafka Topic 动态变化时，Sink 如何不中断发现、路由、写入并恢复。")

    # 2 Agenda
    s = prs.slides.add_slide(contents)
    set_placeholder_text(s, 0, "汇报提纲", 26, True, BLUE)
    agenda = [
        ("01", "研究背景与问题定义", "静态 KafkaSink 与动态 Topic 场景的矛盾"),
        ("02", "需求分析与系统设计", "动态发现、逻辑路由、writer 生命周期、状态提交"),
        ("03", "代码实现与关键机制", "结合 connector/app 模块说明核心类与方法"),
        ("04", "实验设计与结果分析", "发现实时性、吞吐、扩展性、Checkpoint、Exactly-once"),
        ("05", "总结、局限与展望", "原型贡献、适用边界和后续工作"),
    ]
    for i, (num, title, body) in enumerate(agenda):
        y = 1.35 + i * 0.82
        add_text(s, num, 0.95, y, 0.42, 0.23, 13, True, RED)
        add_text(s, title, 1.52, y - 0.02, 2.75, 0.26, 14.5, True, BLUE)
        add_text(s, body, 4.35, y, 4.6, 0.24, 10.7, False, MUTED)
    add_footer(s, 2)
    notes.append("提纲页：强调后面会把论文论述和代码实现对应起来讲。")

    # 3 Background
    s = prs.slides.add_slide(inner)
    add_title(s, "研究背景：实时链路中的 Sink 目标会变化", "Flink + Kafka 是实时数据链路常见组合，但生产端目标通常在作业启动时固化")
    content_panel(s)
    add_image(s, FIG / "kafka_flink_pipeline.png", 0.75, 2.05, 3.85, 2.75)
    numbered_rows(s, [
        ("业务变化", "多租户扩容、按日期/地域拆分、灰度迁移和容灾切换会持续产生新 Topic。"),
        ("静态绑定", "传统 KafkaSink 在构建阶段指定目标 Topic，运行期不会自动感知 Kafka 元数据变化。"),
        ("运维代价", "新增目标往往需要停机、改配置、重新提交作业，削弱实时链路连续性。"),
        ("研究问题", "当可写 Topic 集合变化时，Sink 如何发现变化、更新 writer，并保持状态可恢复？"),
    ], 4.95, 2.08, 4.25, 0.68)
    add_footer(s, 3)
    notes.append("背景页：把问题从“写一个 Topic”提升到“运行期目标集合变化时如何不中断写出”。")

    # 4 Related gap
    s = prs.slides.add_slide(inner)
    add_title(s, "现有方案能力边界", "本文不是重复 KafkaSink，而是在 Sink 侧补齐动态目标管理")
    content_panel(s)
    rows = [
        ["方案", "已有能力", "不足 / 与本文差异"],
        ["静态 KafkaSink", "固定 Topic 写入、投递语义、事务提交", "目标在构建阶段绑定；缺少运行期正则发现和逻辑 route 表"],
        ["Source 正则订阅", "Source 侧可按正则发现 Topic/Partition", "面向消费端 offset 和分区管理；不能直接解决 producer/writer/commit"],
        ["Kafka Connect / CDC", "外部系统接入、任务级配置管理", "关注数据进入 Kafka/Flink，不负责 Flink Sink 内部动态多目标写出"],
        ["本文原型", "Topic 发现 + route 映射 + route 级 writer/state/committable", "目标是本地小规模可运行原型和实验闭环"],
    ]
    add_table(s, rows, 0.78, 2.05, 8.65, 3.65, 8.1)
    add_footer(s, 4)
    notes.append("相关工作页：突出 Sink 侧的额外复杂性，包括 producer 生命周期、writer 状态和事务提交。")

    # 5 Requirements
    s = prs.slides.add_slide(inner)
    add_title(s, "需求分析：功能需求与量化指标", "论文第 3 章将问题拆成四类功能需求和五个可测指标")
    content_panel(s)
    section_label(s, "功能需求", 0.82, 2.05, BLUE)
    bullet_list(s, [
        "动态 Topic 发现：周期性查询 Kafka 元数据，按 stream_pattern 识别 Topic",
        "自动订阅/扩写：发现后维护 activeRouteIds，并创建可写 ClusterWriter",
        "动态路由：业务 routeId 与 Kafka 物理 Topic 解耦，支持运行期 route 更新",
        "数据处理与恢复：序列化、写入、Checkpoint 快照、恢复和提交路径保持一致",
    ], 0.9, 2.48, 4.0, 2.4, 10.7)
    section_label(s, "非功能指标", 5.2, 2.05, BLUE)
    rows = [
        ["指标", "阈值"],
        ["发现实时性", "新增 Topic 首次写入 ≤ 3 s"],
        ["吞吐可接受性", "动态 2-route 较静态下降 ≤ 5%"],
        ["小规模扩展性", "2-route 到 4-route 下降 ≤ 10%"],
        ["Checkpoint 协同", "稳态完成时延 ≤ 100 ms"],
        ["事务路径", "EO ≥ 5 次 Checkpoint 且双 Topic 输出"],
    ]
    add_table(s, rows, 5.22, 2.47, 3.95, 2.65, 8.2)
    add_footer(s, 5)
    notes.append("需求页：说明指标不是生产 SLA，而是针对本地单机原型的验收标准。")

    # 6 Architecture
    s = prs.slides.add_slide(inner)
    add_title(s, "总体设计：动态外壳 + 静态内核", "动态层负责变化管理，底层复用成熟 KafkaSink writer 与提交能力")
    content_panel(s)
    add_image(s, FIG / "dynamic_sink_architecture.png", 0.72, 2.02, 4.95, 3.65)
    numbered_rows(s, [
        ("配置构建层", "DynamicKafkaSinkBuilder 装配 streamPattern、metadata service、serializer 与 delivery guarantee。"),
        ("元数据发现层", "SingleClusterTopicMetadataService 通过 AdminClient.listTopics() 暴露 Topic 集合。"),
        ("运行时路由层", "DynamicKafkaSinkWriter 维护 activeRouteIds、explicitRoutesById 和 clusterWritersByRouteId。"),
        ("状态提交层", "DynamicKafkaSinkWriterState 与 DynamicKafkaCommittable 让 route 信息进入 Checkpoint/Committer。"),
    ], 5.95, 2.08, 3.1, 0.72)
    add_footer(s, 6)
    notes.append("架构页：核心是外层处理动态目标，内层仍是 KafkaSink，降低实现风险。")

    # 7 Project code map
    s = prs.slides.add_slide(inner)
    add_title(s, "代码结构：论文设计在仓库中的落点", "connector 是核心交付，app 是验证载体，docs 保存论文和实验复现材料")
    content_panel(s)
    rows = [
        ["模块/文件", "职责", "论文对应内容"],
        ["connector/.../DynamicKafkaSink.java", "实现 TwoPhaseCommittingStatefulSink，创建 writer/committer/serializer", "第 4 章总体架构与 Sink API 接入"],
        ["connector/.../writer/DynamicKafkaSinkWriter.java", "动态发现、显式路由、writer 创建、快照、预提交", "第 4 章核心流程与 route 生命周期"],
        ["connector/.../metadata/SingleClusterTopicMetadataService.java", "基于 AdminClient 获取 Topic 元数据", "第 4 章元数据发现层"],
        ["connector/.../committer/DynamicKafkaCommitter.java", "按 route 分组提交 KafkaCommittable", "第 4 章 Exactly-once 提交流程"],
        ["app/.../DynamicJob.java / AutoDiscoveryJob.java / RegularJob.java", "广播 route 更新、数据生成、自动发现和静态基线实验", "第 5 章实验载体"],
    ]
    add_table(s, rows, 0.72, 2.02, 8.8, 3.95, 7.4)
    add_footer(s, 7)
    notes.append("代码结构页：让老师知道每个论文设计点在代码中有明确实现。")

    # 8 Metadata discovery code
    s = prs.slides.add_slide(inner)
    add_title(s, "实现一：Topic 元数据发现与正则匹配", "Kafka Topic 被抽象为 KafkaStream，streamId 在当前原型中等价于 Topic 名")
    content_panel(s)
    code_box(s, "SingleClusterTopicMetadataService.getAllStreams()", [
        "return getAdminClient().listTopics().names().get().stream()",
        "    .map(this::createKafkaStream)",
        "    .collect(Collectors.toSet());",
        "",
        "private KafkaStream createKafkaStream(String topic) {",
        "    ClusterMetadata meta = new ClusterMetadata(",
        "        Collections.singleton(topic), properties);",
        "    return new KafkaStream(topic,",
        "        Collections.singletonMap(kafkaClusterId, meta));",
        "}",
    ], 0.8, 2.1, 4.05, 3.25)
    section_label(s, "设计含义", 5.18, 2.12, BLUE)
    bullet_list(s, [
        "元数据事实来源：Kafka AdminClient，而不是 Flink 内部事件通知",
        "发现粒度：Topic 级；不改变 Kafka partition 分配和 producer 发送语义",
        "正则匹配位置：StreamPatternSubscriber 对 KafkaStream.streamId 过滤",
        "发现延迟上界：主要受 discovery_interval_ms 与本地调度影响",
    ], 5.25, 2.55, 3.75, 1.8, 11.2)
    card(s, "论文对应", "第 2 章说明 Source 正则订阅不能直接迁移到 Sink；第 4 章将 Topic 元数据作为动态写入目标来源。", 5.25, 4.62, 3.8, 0.72, BLUE, 9.2)
    add_footer(s, 8)
    notes.append("元数据页：强调动态发现来自 Kafka 管理面轮询，发现的是 Topic，而非 partition。")

    # 9 Writer main flow
    s = prs.slides.add_slide(inner)
    add_title(s, "实现二：Writer 主写入流程", "同一个 write() 入口同时处理 route 更新事件、显式路由数据和自动发现数据")
    content_panel(s)
    code_box(s, "DynamicKafkaSinkWriter.write()", [
        "if (element instanceof DynamicKafkaRouteEvent) {",
        "    if (routeEvent.isRouteUpdate()) {",
        "        applyRouteUpdates(routeEvent.getRouteUpdates());",
        "        return;",
        "    }",
        "    if (routeId != null) {",
        "        writeToExplicitRoute(routeId, element, context);",
        "        return;",
        "    }",
        "}",
        "refreshRouteIfNeeded(false);",
        "for (String routeId : activeRouteIds) {",
        "    clusterWritersByRouteId.get(routeId).writer.write(element, context);",
        "}",
    ], 0.8, 2.08, 4.35, 3.55)
    rows = [
        ["输入类型", "处理方式", "结果"],
        ["route 更新事件", "applyRouteUpdates()", "刷新逻辑 route 表与解析缓存"],
        ["携带 routeId 的数据", "writeToExplicitRoute()", "按逻辑 route 分发到物理 Topic"],
        ["不携带 routeId 的数据", "refreshRouteIfNeeded()", "写入所有 activeRouteIds"],
    ]
    add_table(s, rows, 5.38, 2.12, 3.85, 2.05, 8.0)
    card(s, "关键点", "自动发现路径默认 fan-out 到所有活动 route；显式 route 路径用于业务分流。两条路径最终统一落到 clusterWritersByRouteId。", 5.38, 4.45, 3.85, 0.82, BLUE, 9.2)
    add_footer(s, 9)
    notes.append("Writer 页：这是核心实现。讲清楚三类输入如何分流，以及最终统一走底层 writer。")

    # 10 Dynamic route broadcast
    s = prs.slides.add_slide(inner)
    add_title(s, "实现三：逻辑 route 与物理 Topic 解耦", "应用层用广播状态传播 route_map，connector 内部解析到物理 route")
    content_panel(s)
    code_box(s, "DynamicJob.RouteBroadcastProcessFunction", [
        "routeUpdates = env.fromData(",
        "    DynamicSinkEvent.update(routeMap))",
        "    .broadcast(routeMapStateDesc);",
        "",
        "processElement(value, ctx, out):",
        "  if broadcastState contains value.routeId:",
        "      out.collect(value)",
        "",
        "processBroadcastElement(value, ctx, out):",
        "  update broadcastState(routeId -> destination)",
        "  out.collect(value)  // 继续发送给 Sink",
    ], 0.8, 2.08, 4.2, 3.5)
    code_box(s, "DynamicKafkaSinkWriter.resolveExplicitRoutes()", [
        "topicPattern = Pattern.compile(destination.getTopicPattern());",
        "for (KafkaStream stream : kafkaMetadataService.getAllStreams())",
        "  for (String topic : clusterMetadata.getTopics())",
        "    if topicPattern.matcher(topic).find():",
        "       routeId = logicalRouteId + '|' + clusterId + '|'",
        "               + topic + '|' + destination.bootstrapServers",
        "       createResolvedRoute(routeId, clusterId, topic, destination)",
    ], 5.25, 2.08, 4.05, 3.5)
    add_footer(s, 10)
    notes.append("逻辑路由页：说明业务事件只携带 routeId，物理 Topic 可以通过配置变化切换。")

    # 11 Route refresh lifecycle
    s = prs.slides.add_slide(inner)
    add_title(s, "实现四：route 刷新与 writer 生命周期", "全量重算活动集合，惰性创建缺失 writer，并在非事务路径清理退役 writer")
    content_panel(s)
    code_box(s, "refreshRouteIfNeeded(force)", [
        "if metadataRefreshIntervalMs > 0",
        "   && now < nextMetadataRefreshTimestamp: return",
        "",
        "routes = resolveRoutes();",
        "activeRouteIds = routes.map(routeId).collect(toSet());",
        "nextMetadataRefreshTimestamp = now + interval;",
        "",
        "for route in routes:",
        "  clusterWritersByRouteId.computeIfAbsent(",
        "      route.routeId, ignored -> createClusterWriter(...));",
        "cleanupRetiredWriters();",
    ], 0.8, 2.1, 4.15, 3.38)
    section_label(s, "生命周期策略", 5.25, 2.1, BLUE)
    numbered_rows(s, [
        ("未发现", "Topic 不存在或不匹配正则，不创建 writer。"),
        ("已发现", "元数据刷新命中 Topic，生成 ResolvedRoute。"),
        ("活跃", "首次写入时懒创建 KafkaSink writer，并参与快照/提交。"),
        ("退役", "不再被 active route 或显式 route 引用；非 EO 路径关闭并移除。"),
    ], 5.25, 2.55, 3.7, 0.62, [BLUE, RED])
    card(s, "Exactly-once 取舍", "cleanupRetiredWriters() 在 EXACTLY_ONCE 下直接返回，避免过早关闭仍可能参与未完成事务的 writer。", 5.25, 5.05, 3.75, 0.55, RED, 8.8)
    add_footer(s, 11)
    notes.append("生命周期页：说明当前实现对 EO 路径保守处理，这是合理取舍也是后续工作。")

    # 12 Serialization / underlying sink
    s = prs.slides.add_slide(inner)
    add_title(s, "实现五：每个物理 route 复用一个底层 KafkaSink", "动态层不重写 Kafka producer，而是为 route 克隆 serializer 并创建 KafkaSink writer")
    content_panel(s)
    code_box(s, "createClusterWriter(...)", [
        "KafkaRecordSerializationSchema<IN> clusterSerializer =",
        "    cloneSerializerForRoute(topic);",
        "",
        "TwoPhaseCommittingStatefulSink<IN, KafkaWriterState, KafkaCommittable> clusterSink =",
        "    KafkaSink.<IN>builder()",
        "        .setDeliveryGuarantee(deliveryGuarantee)",
        "        .setKafkaProducerConfig(kafkaProducerConfig)",
        "        .setBootstrapServers(...)",
        "        .setRecordSerializer(clusterSerializer)",
        "        .setTransactionalIdPrefix(routeTransactionalIdPrefix)",
        "        .build();",
        "",
        "writer = recoveredStates.isEmpty()",
        "    ? clusterSink.createWriter(context)",
        "    : clusterSink.restoreWriter(context, recoveredStates);",
    ], 0.8, 2.05, 4.7, 3.8)
    section_label(s, "设计收益", 5.78, 2.1, BLUE)
    bullet_list(s, [
        "复用官方 KafkaSink 的分区、发送、flush、事务提交等成熟逻辑",
        "动态层只负责 route 解析、writer 生命周期和状态/committable 包装",
        "FixedTopicKafkaRecordSerializationSchema 强制改写最终发送 Topic",
        "恢复时可根据 route 级状态重建对应底层 writer",
    ], 5.85, 2.52, 3.45, 2.0, 11.2)
    card(s, "答辩表述", "本文创新点不是重新实现 Kafka producer，而是把动态目标管理纳入 Flink Sink API 的 Writer/State/Committer 阶段。", 5.85, 4.82, 3.45, 0.62, RED, 8.7)
    add_footer(s, 12)
    notes.append("底层 Sink 页：动态外壳 + 静态内核的代码证据。")

    # 13 State and checkpoint
    s = prs.slides.add_slide(inner)
    add_title(s, "实现六：route 级状态进入 Checkpoint", "每个实际写入通道都作为可恢复单元保存")
    content_panel(s)
    code_box(s, "snapshotState(checkpointId)", [
        "refreshRouteIfNeeded(false);",
        "for (ClusterWriter clusterWriter : clusterWritersByRouteId.values()) {",
        "    kafkaStates = clusterWriter.writer.snapshotState(checkpointId);",
        "    states.add(new DynamicKafkaSinkWriterState(",
        "        clusterWriter.routeId,",
        "        clusterWriter.kafkaClusterId,",
        "        clusterWriter.topic,",
        "        clusterWriter.transactionalIdPrefix,",
        "        clusterWriter.kafkaProducerConfig,",
        "        kafkaStates));",
        "}",
    ], 0.8, 2.08, 4.2, 3.5)
    rows = [
        ["状态字段", "作用"],
        ["routeId", "恢复和提交时定位物理 route"],
        ["kafkaClusterId / topic", "重建目标集群与 Topic"],
        ["transactionalIdPrefix", "恢复 EO 事务命名空间"],
        ["kafkaProducerConfig", "恢复 producer 连接配置"],
        ["KafkaWriterState 列表", "复用底层 KafkaSink writer 状态"],
    ]
    add_table(s, rows, 5.25, 2.08, 4.05, 2.9, 8.0)
    card(s, "恢复路径", "DynamicKafkaSink.restoreWriter(...) 将 recoveredState 传入 DynamicKafkaSinkWriter 构造函数，构造函数再按 route 重建 ClusterWriter。", 5.25, 5.18, 4.05, 0.58, BLUE, 8.6)
    add_footer(s, 13)
    notes.append("Checkpoint 页：说明状态字段为什么需要这么多，因为恢复时要重建 route 级 writer。")

    # 14 Commit path
    s = prs.slides.add_slide(inner)
    add_title(s, "实现七：Exactly-once 的 route 级提交路径", "DynamicKafkaCommittable 包装底层 KafkaCommittable，提交器按 route 分组")
    content_panel(s)
    code_box(s, "prepareCommit()", [
        "for (ClusterWriter clusterWriter : clusterWritersByRouteId.values())",
        "  for (KafkaCommittable kc : clusterWriter.writer.prepareCommit())",
        "    committables.add(new DynamicKafkaCommittable(",
        "        clusterWriter.routeId,",
        "        clusterWriter.kafkaClusterId,",
        "        clusterWriter.topic,",
        "        clusterWriter.transactionalIdPrefix,",
        "        clusterWriter.kafkaProducerConfig,",
        "        kc));",
    ], 0.8, 2.08, 4.15, 2.7)
    code_box(s, "DynamicKafkaCommitter.commit()", [
        "if deliveryGuarantee != EXACTLY_ONCE: return;",
        "",
        "requestsByRouteId = groupBy(request.getCommittable().getRouteId());",
        "for each routeId:",
        "  kafkaCommitter = committersByRouteId.computeIfAbsent(routeId,",
        "      new KafkaCommitter(routeConfig, transactionalIdPrefix, ...));",
        "  kafkaCommitter.commit(delegatingKafkaRequests);",
    ], 5.2, 2.08, 4.15, 2.7)
    card(s, "事务隔离", "buildTransactionalIdPrefix(routeId) = 全局 transactionalIdPrefix + routeId.hashCode()，避免多个物理 route 共享事务前缀。", 0.85, 5.05, 8.4, 0.55, RED, 8.6)
    add_footer(s, 14)
    notes.append("提交页：强调按 route 分组，避免不同目标的事务上下文混在一起。")

    # 15 Experiment design
    s = prs.slides.add_slide(inner)
    add_title(s, "实验设计：需求、作业与指标一一对应", "app 模块提供 AutoDiscoveryJob、DynamicJob、RegularJob 三类验证载体")
    content_panel(s)
    rows = [
        ["需求/指标", "实验作业", "配置与观测"],
        ["动态 Topic 发现", "AutoDiscoveryJob", "stream_pattern='^exp-auto-.*$'；运行中创建 Topic；轮询 offset"],
        ["显式 route 分发", "DynamicJob", "route-a/b → exp-dynamic-c/d；统计两个目标 Topic offset"],
        ["静态基线对比", "RegularJob", "固定写入 exp-regular-2；同样运行约 30 s"],
        ["小规模扩展性", "DynamicJob", "2-route 与 4-route 配置对照"],
        ["Checkpoint 协同", "DynamicJob", "checkpoint-interval-ms=5000；观察成功次数和完成时延"],
        ["事务路径", "DynamicJob + EXACTLY_ONCE", "指定 transactional-id-prefix；观察 committed 输出"],
    ]
    add_table(s, rows, 0.72, 2.05, 8.75, 3.85, 7.6)
    add_text(s, "实验边界：本地 Kafka、parallelism=1、小规模 route、短时间窗口；结论用于原型验证，不外推为生产性能。", 0.88, 6.08, 8.25, 0.28, 9.8, False, MUTED)
    add_footer(s, 15)
    notes.append("实验设计页：强调不是泛泛测试，而是每个需求都有对应实验。")

    # 16 Discovery results
    s = prs.slides.add_slide(inner)
    add_title(s, "实验结果一：动态发现实时性", "新增匹配 Topic 可在不重启作业的情况下进入写入路径")
    content_panel(s)
    metric(s, "6", "重复样本数", 0.9, 2.12, 1.45, 0.78, BLUE)
    metric(s, "1.530 s", "平均首次写入延迟", 2.55, 2.12, 1.7, 0.78, BLUE)
    metric(s, "2.418 s", "最大首次写入延迟", 4.45, 2.12, 1.7, 0.78, RED)
    metric(s, "≤ 3 s", "需求阈值", 6.35, 2.12, 1.45, 0.78, MUTED)
    rows = [
        ["轮次", "Topic 编号", "首次写入延迟 / s", "首次观测 offset"],
        ["第 1 轮", "1", "1.077", "14"],
        ["第 1 轮", "2", "2.412", "23"],
        ["第 1 轮", "3", "1.095", "8"],
        ["第 2 轮", "1", "1.103", "14"],
        ["第 2 轮", "2", "2.418", "24"],
        ["第 2 轮", "3", "1.074", "10"],
    ]
    add_table(s, rows, 0.9, 3.22, 5.7, 2.25, 7.8)
    bullet_list(s, [
        "结果落在 1.074-2.418 s，与 discovery_interval_ms=2000 的轮询机制吻合",
        "最大值低于 3 s 阈值，验证自动发现路径可用",
        "延迟主因是等待下一次元数据刷新，而非 Kafka 写入本身",
    ], 6.82, 3.25, 2.25, 1.8, 9.5)
    add_footer(s, 16)
    notes.append("发现结果页：讲最大值低于 3 秒，且结果符合轮询刷新机制。")

    # 17 Throughput results
    s = prs.slides.add_slide(inner)
    add_title(s, "实验结果二：吞吐可接受性与扩展性", "动态 2-route 与静态基线吞吐接近，4-route 未出现明显下降")
    content_panel(s)
    rows = [
        ["方案", "语义", "route 数", "输出总条数", "近似吞吐"],
        ["静态基线", "AT_LEAST_ONCE", "1", "5776", "192.5 条/s"],
        ["动态写入", "AT_LEAST_ONCE", "2", "5742", "191.4 条/s"],
        ["动态写入", "AT_LEAST_ONCE", "4", "5763", "192.1 条/s"],
        ["动态写入", "EXACTLY_ONCE", "2", "2092", "95.1 条/s"],
    ]
    add_table(s, rows, 0.8, 2.08, 5.45, 2.35, 8.5)
    metric(s, "0.59%", "2-route 较静态基线下降", 6.55, 2.1, 2.1, 0.82, RED)
    metric(s, "无下降", "4-route 相对 2-route", 6.55, 3.12, 2.1, 0.82, BLUE)
    metric(s, "95.1/s", "EO 路径吞吐", 6.55, 4.14, 2.1, 0.82, MUTED)
    bullet_list(s, [
        "route-a / route-b 输出：2812 / 2930，约 49.0% / 51.0%，分布均衡",
        "动态管理在当前本地小规模场景下不是主要瓶颈",
        "EXACTLY_ONCE 吞吐约为 AT_LEAST_ONCE 的一半，成本来自事务与 Checkpoint 提交协同",
    ], 0.9, 4.82, 8.2, 0.9, 9.8)
    add_footer(s, 17)
    notes.append("吞吐结果页：重点是动态开销 0.59%，4-route 没有下降，EO 成本明显但可运行。")

    # 18 Checkpoint and EO results
    s = prs.slides.add_slide(inner)
    add_title(s, "实验结果三：Checkpoint 协同与事务路径", "route 级状态和 committable 能参与 Flink 原生容错/提交链路")
    content_panel(s)
    section_label(s, "Checkpoint 协同", 0.85, 2.08, BLUE)
    metric(s, "5 次", "连续成功 Checkpoint", 0.9, 2.55, 1.55, 0.78, BLUE)
    metric(s, "18.2 ms", "稳态平均完成时延", 2.65, 2.55, 1.7, 0.78, RED)
    metric(s, "≤100 ms", "需求阈值", 4.55, 2.55, 1.45, 0.78, MUTED)
    bullet_list(s, [
        "首次 Checkpoint 476 ms，后续 11/22/19/21 ms",
        "source、广播控制流、dynamic sink 均收到 checkpoint completed 日志",
    ], 0.9, 3.58, 4.8, 0.8, 9.8)
    section_label(s, "Exactly-once 可运行性", 0.85, 4.62, RED)
    metric(s, "2092", "EO 输出总条数", 0.9, 5.05, 1.55, 0.78, RED)
    metric(s, "1084 / 1008", "两个目标 Topic 输出", 2.65, 5.05, 1.95, 0.78, BLUE)
    metric(s, "5 次", "EO 成功 Checkpoint", 4.85, 5.05, 1.35, 0.78, BLUE)
    card(s, "结论", "满足“至少 5 次 Checkpoint 且两个目标 Topic 均产生 committed 输出”的判定标准；但尚不等同于复杂故障注入下的端到端精准一次证明。", 6.45, 2.55, 2.6, 3.25, RED, 9.3)
    add_footer(s, 18)
    notes.append("Checkpoint/EO 页：既讲成功，也诚实说明没有做复杂故障注入。")

    # 19 Comparison
    s = prs.slides.add_slide(inner)
    add_title(s, "对比分析：动态方案的价值与成本", "相较静态 KafkaSink，本文以很小本地吞吐代价换取无重启目标适配")
    content_panel(s)
    add_image(s, FIG / "offline_update_architecture.png", 0.75, 2.05, 3.55, 2.45)
    rows = [
        ["维度", "静态 KafkaSink", "本文 Dynamic Sink"],
        ["目标绑定", "构建阶段指定固定 Topic", "运行期正则发现 + 逻辑 route 映射"],
        ["变更方式", "通常停机修改配置并重提作业", "通过 metadata refresh / route update 接纳变化"],
        ["状态粒度", "固定目标 writer 状态", "route 级 writer state + committable"],
        ["运行成本", "实现简单、开销低", "需要元数据轮询、多 writer 管理和提交包装"],
    ]
    add_table(s, rows, 4.65, 2.05, 4.6, 2.75, 7.6)
    bullet_list(s, [
        "适用场景：Topic 集合长期稳定时静态 KafkaSink 更简单；Topic 频繁增删或迁移时动态 Sink 更有价值",
        "本文方案保留 KafkaSink 的成熟写入能力，同时补充运行期目标管理",
    ], 0.88, 5.1, 8.2, 0.72, 10.2)
    add_footer(s, 19)
    notes.append("对比页：说明不是否定静态 KafkaSink，而是面向不同场景。")

    # 20 Summary
    s = prs.slides.add_slide(inner)
    add_title(s, "总结：论文与代码形成闭环", "当前原型完成了动态发现、路由解耦、route 级状态提交和本地实验验证")
    content_panel(s)
    card(s, "贡献一：Sink 侧动态发现", "实现 Kafka 元数据轮询 + stream_pattern 匹配，新增 Topic 最大 2.418 s 内开始写入。", 0.85, 2.08, 2.62, 1.05, RED, 9.3)
    card(s, "贡献二：逻辑路由解耦", "通过 DynamicSinkEvent / KafkaRouteDestination 将业务 routeId 与物理 Topic 分离，支持运行期更新。", 3.65, 2.08, 2.62, 1.05, BLUE, 9.3)
    card(s, "贡献三：route 级容错提交", "DynamicKafkaSinkWriterState 与 DynamicKafkaCommittable 让动态目标进入 Checkpoint 和 Committer。", 6.45, 2.08, 2.62, 1.05, BLUE, 9.3)
    section_label(s, "实验结论", 0.9, 3.6, BLUE)
    bullet_list(s, [
        "吞吐：动态 2-route 相对静态基线下降约 0.59%，低于 5% 阈值",
        "扩展：2-route 到 4-route 未出现吞吐下降，满足小规模扩展性指标",
        "Checkpoint：稳态时延约 18.2 ms，低于 100 ms 阈值",
        "事务：EXACTLY_ONCE 本地小规模路径完成 5 次 Checkpoint 且双 Topic 输出",
    ], 0.95, 4.02, 7.95, 1.1, 10.3)
    add_footer(s, 20)
    notes.append("总结页：把贡献和实验结论压缩成一页，答辩收束用。")

    # 21 Limitations
    s = prs.slides.add_slide(inner)
    add_title(s, "不足与展望", "当前成果是可运行原型，后续需要更完整的生产化验证")
    content_panel(s)
    numbered_rows(s, [
        ("元数据访问优化", "当前依赖周期性 AdminClient.listTopics()；后续可引入缓存、增量发现和多集群元数据服务。"),
        ("资源生命周期", "route 数增长会带来 writer/producer 资源压力；EO 路径下退役 writer 安全回收仍需完善。"),
        ("一致性验证", "需要故障注入、任务恢复、事务超时、read_committed 消费校验等端到端测试。"),
        ("性能评估", "扩展到更高并行度、更多 route、更长运行时间和真实数据流，区分 Kafka 瓶颈与动态管理开销。"),
    ], 0.9, 2.08, 7.85, 0.82)
    card(s, "答辩时的边界表述", "本文证明的是本地单机、小规模 route 场景下动态 Kafka Sink 机制可行；生产级 SLA 和复杂故障场景属于后续工程化工作。", 0.95, 5.55, 8.1, 0.62, RED, 9.2)
    add_footer(s, 21)
    notes.append("不足页：主动说明边界，避免过度承诺。")

    # 22 End
    s = prs.slides.add_slide(end)
    set_placeholder_text(s, 0, "谢谢各位老师\n欢迎批评指正", 30, True, BLUE)
    notes.append("结束页：准备回答问题。")

    prs.save(OUT)
    NOTES.write_text("\n\n".join(f"## 第 {i + 1} 页\n{n}" for i, n in enumerate(notes)), encoding="utf-8")
    print(OUT)
    print(NOTES)


if __name__ == "__main__":
    make_deck()
