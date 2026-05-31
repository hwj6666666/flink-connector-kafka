from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "基于Flink的动态Kafka-Topic发现与订阅机制研究_毕设答辩.pptx"
NOTES = ROOT / "docs" / "基于Flink的动态Kafka-Topic发现与订阅机制研究_答辩讲稿.md"
FIG = ROOT / "docs" / "SJTUThesis" / "figures"

W, H = Inches(13.333), Inches(7.5)
RED = RGBColor(180, 32, 37)
BLUE = RGBColor(20, 72, 121)
INK = RGBColor(34, 39, 46)
MUTED = RGBColor(93, 102, 112)
LIGHT = RGBColor(246, 248, 250)
LINE = RGBColor(218, 222, 227)
GREEN = RGBColor(31, 139, 86)
ORANGE = RGBColor(206, 112, 35)
WHITE = RGBColor(255, 255, 255)


def add_textbox(slide, text, x, y, w, h, size=22, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "PingFang SC"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, 0.65, 0.38, 9.9, 0.48, 26, True, INK)
    if subtitle:
        add_textbox(slide, subtitle, 0.68, 0.9, 9.4, 0.28, 10, False, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.22), Inches(1.25), Inches(0.045))
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()


def add_footer(slide, page):
    add_textbox(slide, "基于 Flink 的动态 Kafka Topic 发现与订阅机制研究", 0.65, 7.05, 6.2, 0.18, 8.5, False, MUTED)
    add_textbox(slide, f"{page:02d}", 12.25, 7.04, 0.45, 0.18, 8.5, False, MUTED, PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, h, size=17, color=INK, gap=0.1):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "PingFang SC"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
        p.space_after = Pt(gap * 12)
    return box


def card(slide, title, body, x, y, w, h, accent=BLUE, body_size=13.5):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.08
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_textbox(slide, title, x + 0.22, y + 0.16, w - 0.35, 0.3, 14.5, True, accent)
    add_textbox(slide, body, x + 0.22, y + 0.58, w - 0.35, h - 0.72, body_size, False, INK)


def metric(slide, value, label, x, y, w=2.3, h=1.0, color=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(250, 251, 252)
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.08
    add_textbox(slide, value, x + 0.12, y + 0.13, w - 0.24, 0.34, 22, True, color, PP_ALIGN.CENTER)
    add_textbox(slide, label, x + 0.12, y + 0.58, w - 0.24, 0.25, 9.5, False, MUTED, PP_ALIGN.CENTER)


def add_image_fit(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_table(slide, rows, x, y, w, h, font_size=11):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(242, 245, 248) if r == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = "PingFang SC"
                p.font.size = Pt(font_size)
                p.font.bold = r == 0
                p.font.color.rgb = INK if r else BLUE
    return table


def make_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    notes = []

    # 1
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(248, 249, 251)
    bg.line.fill.background()
    add_textbox(s, "基于 Flink 的动态 Kafka Topic\n发现与订阅机制研究", 0.72, 1.36, 8.2, 1.45, 34, True, INK)
    add_textbox(s, "本科毕业设计答辩", 0.75, 3.0, 3.2, 0.35, 18, False, RED)
    add_textbox(s, "答辩人：胡文杰    指导教师：任锐\n上海交通大学    2026 年 5 月", 0.75, 5.58, 5.9, 0.58, 14, False, MUTED)
    add_image_fit(s, FIG / "dynamic_sink_architecture.png", 7.1, 1.0, 5.45, 4.55)
    notes.append("开场：介绍题目和研究对象。强调课题关注 Flink 写出端，即 Sink 侧在 Kafka Topic 动态变化时如何不中断地发现、路由和提交。")

    # 2
    s = prs.slides.add_slide(blank)
    add_title(s, "汇报提纲", "围绕“为什么做、怎么做、做得如何、还差什么”展开")
    sections = [
        ("01", "研究背景与问题", "静态 Kafka Sink 难以适应运行期 Topic 变化"),
        ("02", "需求与总体设计", "动态发现、逻辑路由、writer 管理、Checkpoint"),
        ("03", "核心实现机制", "动态外壳复用静态 KafkaSink 内核"),
        ("04", "实验验证结果", "发现实时性、吞吐、扩展性、事务路径"),
        ("05", "总结与展望", "原型价值、局限与后续优化方向"),
    ]
    for i, (num, title, body) in enumerate(sections):
        y = 1.45 + i * 0.92
        add_textbox(s, num, 0.88, y, 0.45, 0.28, 14, True, RED)
        add_textbox(s, title, 1.42, y - 0.02, 2.8, 0.3, 17, True, INK)
        add_textbox(s, body, 4.15, y, 6.7, 0.28, 13, False, MUTED)
    add_footer(s, 2)
    notes.append("提纲页：后面按背景、设计、实现、实验、总结五部分讲。答辩时可控制在每部分 1-2 分钟。")

    # 3
    s = prs.slides.add_slide(blank)
    add_title(s, "研究背景", "实时链路中的下游 Kafka Topic 并不总是稳定不变")
    add_image_fit(s, FIG / "kafka_flink_pipeline.png", 0.8, 1.45, 5.6, 3.55)
    card(s, "场景变化", "多租户、按日期/地域拆分、灰度迁移、容灾切换都会带来 Topic 集合变化。", 6.75, 1.45, 5.3, 1.05, RED)
    card(s, "静态 Sink 的问题", "传统 KafkaSink 在作业构建阶段绑定目标 Topic；目标变化通常意味着改配置、停作业、重新提交。", 6.75, 2.78, 5.3, 1.05, ORANGE)
    card(s, "本文关注点", "不是“如何写入一个 Topic”，而是“目标集合变化时，Sink 如何发现、路由、写出并可恢复”。", 6.75, 4.11, 5.3, 1.05, BLUE)
    add_footer(s, 3)
    notes.append("背景页：讲清楚静态 KafkaSink 在拓扑稳定时足够好，但在 Topic 频繁增删或迁移时，停机变更会破坏实时链路连续性。")

    # 4
    s = prs.slides.add_slide(blank)
    add_title(s, "问题定义与研究目标", "把目标变化纳入 Flink Sink 的运行时与容错链路")
    bullet_list(s, [
        "动态 Topic 发现：按正则周期查询 Kafka 元数据，识别新增可写目标",
        "自动订阅/扩写：发现目标后维护 active routes，并懒创建底层 writer",
        "动态路由：业务 route id 与物理 Kafka Topic 解耦，支持运行期 route 更新",
        "状态与提交：route 级 writer 状态进入 Checkpoint，committable 按 route 分组提交",
    ], 0.9, 1.55, 6.0, 2.55, 17)
    rows = [
        ["指标", "目标"],
        ["发现实时性", "新增匹配 Topic 首次写入 ≤ 3 s"],
        ["吞吐可接受性", "动态 2-route 相对静态基线下降 ≤ 5%"],
        ["小规模扩展性", "2-route 到 4-route 下降 ≤ 10%"],
        ["Checkpoint 协同", "稳态完成时延 ≤ 100 ms"],
        ["事务路径", "EXACTLY_ONCE 至少 5 次 Checkpoint 且双 Topic 输出"],
    ]
    add_table(s, rows, 7.1, 1.45, 5.3, 3.9, 10.5)
    add_footer(s, 4)
    notes.append("问题定义页：强调本文同时覆盖功能需求和可测指标。指标来自本地单机原型目标，不外推为生产 SLA。")

    # 5
    s = prs.slides.add_slide(blank)
    add_title(s, "项目代码结构评审", "当前仓库由 connector 原型、验证 app 与论文材料组成")
    card(s, "connector 模块", "核心交付：DynamicKafkaSink、DynamicKafkaSinkWriter、元数据服务、route 状态、committable 与 committer。", 0.85, 1.45, 3.8, 1.55, BLUE)
    card(s, "app 模块", "实验载体：DynamicJob、AutoDiscoveryJob、RegularJob；用于生成数据、广播 route 更新并验证写入。", 4.85, 1.45, 3.8, 1.55, GREEN)
    card(s, "docs 与配置", "论文、实验 runbook、架构图、多个 config.experiment*.yaml；具备较完整复现实验材料。", 8.85, 1.45, 3.8, 1.55, RED)
    bullet_list(s, [
        "技术栈：Java 17、Flink 2.1.0、Kafka clients 4.2.0、Maven 多模块工程",
        "实现思路清晰：动态能力集中在 connector，app 仅作验证载体",
        "可答辩亮点：不是手写 producer，而是接入 Flink 统一 Sink API 与 Checkpoint/Committer 机制",
        "需如实说明：实验为本地小规模原型，Exactly-once 仍缺少复杂故障注入验证",
    ], 1.0, 3.55, 11.1, 2.45, 16)
    add_footer(s, 5)
    notes.append("项目评审页：给老师一个代码层面的可信感。说明 connector/app/docs 分层明确，也主动说明当前局限。")

    # 6
    s = prs.slides.add_slide(blank)
    add_title(s, "总体架构设计", "“动态外壳 + 静态内核”：外层管理变化，内层复用 KafkaSink")
    add_image_fit(s, FIG / "dynamic_sink_architecture.png", 0.75, 1.35, 7.25, 4.8)
    card(s, "四层职责", "配置构建层 → 元数据发现层 → 运行时路由层 → 状态与提交层", 8.35, 1.52, 3.95, 0.92, BLUE)
    card(s, "关键取舍", "不绕开 Flink Sink API；每个物理 route 下沉为普通 KafkaSink writer，动态层只做解析、生命周期和包装。", 8.35, 2.72, 3.95, 1.25, RED)
    card(s, "收益", "降低原型复杂度，同时让动态 route 能进入原生 Checkpoint 与事务提交流程。", 8.35, 4.25, 3.95, 1.0, GREEN)
    add_footer(s, 6)
    notes.append("架构页：核心句是动态外壳复用静态内核。外层解决 Topic/route 变化，内层复用 KafkaSink 成熟的写入和提交能力。")

    # 7
    s = prs.slides.add_slide(blank)
    add_title(s, "核心流程：发现、路由与写入", "Writer 同时处理数据事件和 route 更新事件")
    add_image_fit(s, FIG / "dynamic_sink_call_graph.png", 0.72, 1.35, 6.5, 4.75)
    rows = [
        ["输入", "处理逻辑", "结果"],
        ["route 更新事件", "更新 explicitRoutesById，清理解析缓存", "后续数据使用新路由"],
        ["携带 routeId", "解析逻辑 route 到物理 Topic", "写入对应 writer"],
        ["不携带 routeId", "按周期刷新元数据，遍历 activeRouteIds", "fan-out 到当前活动目标"],
    ]
    add_table(s, rows, 7.45, 1.55, 5.05, 2.45, 10)
    bullet_list(s, [
        "Topic 发现：AdminClient.listTopics() + stream_pattern 正则匹配",
        "writer 创建：computeIfAbsent 懒创建，避免刷新时重建全部 producer",
        "route id：包含 cluster、topic、bootstrapServers，支撑状态恢复与提交分组",
    ], 7.55, 4.32, 4.85, 1.25, 12.8)
    add_footer(s, 7)
    notes.append("核心流程页：用三类输入讲 Writer 逻辑。自动发现路径适合所有匹配 Topic 都接收数据；显式 route 适合业务分流。")

    # 8
    s = prs.slides.add_slide(blank)
    add_title(s, "状态与 Exactly-once 提交", "动态 route 不能游离在 Flink 容错体系之外")
    card(s, "route 级状态", "DynamicKafkaSinkWriterState 记录 routeId、cluster、topic、事务前缀、producer 配置和底层 KafkaWriterState。", 0.85, 1.4, 5.45, 1.25, BLUE)
    card(s, "route 级 committable", "prepareCommit() 将底层 KafkaCommittable 包装为 DynamicKafkaCommittable，提交器按 route 分组委托。", 0.85, 2.95, 5.45, 1.25, RED)
    card(s, "事务前缀隔离", "EXACTLY_ONCE 模式下为每个物理 route 派生事务前缀，避免不同目标混用事务上下文。", 0.85, 4.5, 5.45, 1.25, GREEN)
    rows = [
        ["阶段", "设计要点"],
        ["构建", "要求提供 transactional-id-prefix"],
        ["写入", "按物理 route 派生事务前缀"],
        ["预提交", "包装底层 committable，携带 route 元数据"],
        ["提交", "DynamicKafkaCommitter 按 route 分发提交"],
        ["回收", "非事务路径清理退役 writer；EO 路径保守保留"],
    ]
    add_table(s, rows, 7.0, 1.45, 5.4, 4.1, 10.5)
    add_footer(s, 8)
    notes.append("状态与提交页：说明这是和普通 KafkaSink 最大的差异之一。动态 route 必须被快照和提交感知，否则故障恢复会丢失运行期目标。")

    # 9
    s = prs.slides.add_slide(blank)
    add_title(s, "实验设计", "用本地单机原型验证机制闭环")
    rows = [
        ["需求", "实验", "指标/判定"],
        ["动态 Topic 发现", "AutoDiscoveryJob 运行中新增 Topic", "首次写入延迟 ≤ 3 s"],
        ["动态路由", "DynamicJob + route_map", "route-a/b 正确分发且均衡"],
        ["吞吐可接受性", "DynamicJob vs RegularJob", "下降 ≤ 5%"],
        ["小规模扩展性", "2-route vs 4-route", "下降 ≤ 10%"],
        ["Checkpoint", "开启 5s checkpoint", "稳态时延 ≤ 100 ms"],
        ["事务路径", "EXACTLY_ONCE 运行", "5 次 checkpoint 且双 Topic 输出"],
    ]
    add_table(s, rows, 0.85, 1.45, 11.8, 4.25, 11)
    bullet_list(s, [
        "实验边界：本地 Kafka、parallelism=1、小规模 route、约 30 秒窗口",
        "统计方式：Kafka topic offset 近似输出条数；日志观察 checkpoint 与事务提交",
    ], 1.0, 6.0, 10.8, 0.6, 14)
    add_footer(s, 9)
    notes.append("实验设计页：主动限定实验边界，避免老师追问时显得过度宣称。强调实验目标是原型可行性验证。")

    # 10
    s = prs.slides.add_slide(blank)
    add_title(s, "实验结果：实时性与 Checkpoint", "发现延迟由 discovery interval 主导，route 状态快照开销较低")
    metric(s, "2.418 s", "新增 Topic 首次写入最大值", 0.85, 1.55, 2.55, 1.15, RED)
    metric(s, "1.530 s", "发现延迟平均值", 3.65, 1.55, 2.35, 1.15, BLUE)
    metric(s, "18.2 ms", "稳态 Checkpoint 平均时延", 6.25, 1.55, 2.75, 1.15, GREEN)
    metric(s, "5 次", "连续成功 Checkpoint", 9.25, 1.55, 2.25, 1.15, ORANGE)
    rows = [
        ["样本", "结果"],
        ["发现实时性", "6 个样本，1.074-2.418 s，满足 ≤ 3 s"],
        ["Checkpoint 协同", "5 次成功；首次 476 ms，后续稳态均值 18.2 ms"],
        ["结论", "新增 Topic 可在秒级接入，route 级状态快照未明显阻塞本地小规模作业"],
    ]
    add_table(s, rows, 1.15, 3.3, 10.75, 2.15, 12)
    add_footer(s, 10)
    notes.append("实时性结果页：发现最大 2.418 秒低于 3 秒；Checkpoint 首次较高，稳态 18.2ms，说明动态状态快照开销可接受。")

    # 11
    s = prs.slides.add_slide(blank)
    add_title(s, "实验结果：吞吐与扩展性", "动态管理层在小规模 AT_LEAST_ONCE 场景下开销较小")
    rows = [
        ["方案", "语义", "route 数", "输出总条数", "近似吞吐"],
        ["静态基线", "AT_LEAST_ONCE", "1", "5776", "192.5 条/s"],
        ["动态写入", "AT_LEAST_ONCE", "2", "5742", "191.4 条/s"],
        ["动态写入", "AT_LEAST_ONCE", "4", "5763", "192.1 条/s"],
        ["动态写入", "EXACTLY_ONCE", "2", "2092", "95.1 条/s"],
    ]
    add_table(s, rows, 0.85, 1.45, 7.2, 3.1, 11.5)
    metric(s, "0.59%", "2-route 相对静态基线吞吐下降", 8.45, 1.55, 3.0, 1.12, GREEN)
    metric(s, "≈0", "4-route 未出现明显下降", 8.45, 2.92, 3.0, 1.12, BLUE)
    metric(s, "约 1/2", "EXACTLY_ONCE 吞吐相对 ALO 下降", 8.45, 4.29, 3.0, 1.12, ORANGE)
    add_textbox(s, "解释：在当前发送速率与小规模 route 下，动态 route 管理不是主要瓶颈；事务初始化、Checkpoint 协同与提交路径成本更高。", 1.0, 5.35, 10.9, 0.55, 14, False, MUTED)
    add_footer(s, 11)
    notes.append("吞吐结果页：重点讲 0.59% 和 4-route 不下降；Exactly-once 吞吐降低是符合一致性语义成本的。")

    # 12
    s = prs.slides.add_slide(blank)
    add_title(s, "对比分析", "动态方案以较小本地吞吐代价换取无重启适配能力")
    add_image_fit(s, FIG / "offline_update_architecture.png", 0.85, 1.42, 5.15, 3.28)
    card(s, "静态 KafkaSink", "优点：简单成熟、单目标开销低。\n不足：目标变化时通常要停机、改配置、重提作业。", 6.35, 1.42, 5.55, 1.35, ORANGE)
    card(s, "本文 Dynamic Sink", "优点：运行期发现 Topic，route 与物理 Topic 解耦，状态与提交按 route 纳入 Flink 机制。", 6.35, 3.02, 5.55, 1.35, BLUE)
    card(s, "与 Source 正则订阅不同", "Sink 侧要额外处理 producer 生命周期、topic 绑定、writer 状态、事务与 committer。", 6.35, 4.62, 5.55, 1.1, RED)
    add_footer(s, 12)
    notes.append("对比页：不要说静态方案不好，而是说适用场景不同。本文价值在 Topic 变化时无需停机，以及把动态目标纳入 Flink 容错链路。")

    # 13
    s = prs.slides.add_slide(blank)
    add_title(s, "总结与创新点", "原型已形成需求、设计、实现、实验之间的闭环")
    card(s, "1. Sink 侧动态发现", "基于 Kafka 元数据轮询与正则匹配，在不重启 Flink 作业的前提下接纳新增 Topic。", 0.85, 1.45, 3.75, 1.35, RED)
    card(s, "2. 逻辑路由解耦", "通过 DynamicSinkEvent 与 KafkaRouteDestination，把业务 route id 与物理 Topic 分离。", 4.8, 1.45, 3.75, 1.35, BLUE)
    card(s, "3. route 级状态提交", "以 route 为粒度包装 writer state 与 committable，让动态目标进入 Checkpoint/Committer。", 8.75, 1.45, 3.75, 1.35, GREEN)
    bullet_list(s, [
        "功能验证：新增 Topic 可发现并写入，显式 route 能正确分发到目标 Topic",
        "性能验证：发现最大 2.418 s；动态 2-route 吞吐下降约 0.59%；稳态 Checkpoint 约 18.2 ms",
        "一致性验证：EXACTLY_ONCE 本地小规模路径完成 5 次 Checkpoint 且双 Topic committed 输出",
    ], 1.0, 3.45, 11.1, 1.55, 16)
    add_footer(s, 13)
    notes.append("总结页：把创新点和实验结果捆绑起来讲。核心贡献不是某个单点，而是动态发现、路由解耦、route 级状态提交的组合。")

    # 14
    s = prs.slides.add_slide(blank)
    add_title(s, "不足与展望", "当前成果是可运行原型，后续需要面向生产化继续打磨")
    card(s, "元数据访问优化", "当前依赖周期性 AdminClient.listTopics()，后续可探索缓存、增量发现和多集群元数据服务。", 0.85, 1.45, 5.6, 1.15, BLUE)
    card(s, "资源与生命周期", "route 数增长会带来 writer/producer 资源压力；EXACTLY_ONCE 下退役 writer 安全回收仍需完善。", 0.85, 2.9, 5.6, 1.15, RED)
    card(s, "更严格一致性验证", "需要加入故障注入、任务恢复、事务超时、read_committed 消费校验等端到端测试。", 6.85, 1.45, 5.6, 1.15, ORANGE)
    card(s, "更完整性能评估", "扩展到更高并行度、更长运行时间、更多 route 和真实数据流，区分 Kafka 瓶颈与动态管理开销。", 6.85, 2.9, 5.6, 1.15, GREEN)
    add_textbox(s, "谢谢各位老师，欢迎批评指正。", 0.95, 5.52, 10.8, 0.55, 24, True, INK, PP_ALIGN.CENTER)
    add_footer(s, 14)
    notes.append("结束页：主动承认局限。可以强调本文完成的是机制原型和本地闭环验证，生产级还需要更长期、更复杂的测试。")

    prs.save(OUT)
    NOTES.write_text("\n\n".join(f"## 第 {i+1} 页\n{n}" for i, n in enumerate(notes)), encoding="utf-8")
    return OUT, NOTES


if __name__ == "__main__":
    out, notes = make_deck()
    print(out)
    print(notes)
